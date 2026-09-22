"""docxtpl.accel must be invisible except for speed."""
import io

from docx import Document
from docx.oxml.xmlchemy import BaseOxmlElement
from lxml import etree
import pytest

import docxtpl
from docxtpl import accel


@pytest.fixture
def accelerated():
    accel.enable()
    yield
    accel.disable()


def observe(path):
    """A broad sweep of python-docx APIs that are implemented with XPath."""
    doc = Document(str(path))
    seen = []
    for table in doc.tables[:12]:
        for row in table.rows[:40]:
            for cell in row.cells:
                seen.append((cell.text, cell.grid_span, [r.text for p in cell.paragraphs for r in p.runs][:4]))
    seen.append([(p.text, p.style.name if p.style is not None else None) for p in doc.paragraphs[:200]])
    seen.append([(s.start_type, s.page_width, [p.text for p in s.header.paragraphs]) for s in doc.sections])
    seen.append(sorted(s.style_id for s in doc.styles))
    seen.append([(i.width, i.height) for i in doc.inline_shapes])
    doc.add_paragraph("tail").add_run("x").bold = True
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(1, 2)).text = "merged"
    stream = io.BytesIO()
    doc.save(stream)
    return seen, etree.tostring(doc.element, method="c14n")


def test_off_by_default():
    assert not accel.is_enabled()
    assert BaseOxmlElement.xpath is accel._original_xpath


def test_same_observations_and_output(corpus):
    sample = [p for p in corpus if p.name in ("template-prepared.docx", "expected-library.docx")][:8]
    expected = [observe(p) for p in sample]
    accel.enable()
    try:
        assert [observe(p) for p in sample] == expected
    finally:
        accel.disable()


def test_errors_and_result_types_are_canonical(accelerated):
    body = Document().element.body
    with pytest.raises(etree.XPathEvalError) as fast:
        body.xpath("unknown:p")
    accel.disable()
    with pytest.raises(etree.XPathEvalError) as plain:
        body.xpath("unknown:p")
    assert str(fast.value) == str(plain.value)
    expressions = ("count(./w:sectPr)", "string(./w:sectPr/w:pgSz/@w:w)", "./w:sectPr/w:pgSz/@w:w",
                   "./w:sectPr", "boolean(./w:p)", "./w:sectPr/w:pgSz/@w:w = '12240'")

    def describe(value):
        if isinstance(value, list):
            return [describe(item) for item in value]
        parent = value.getparent() if hasattr(value, "getparent") else None
        return (type(value).__name__, str(value) if not hasattr(value, "tag") else value.tag,
                getattr(value, "is_attribute", None), None if parent is None else parent.tag)

    plain = [describe(body.xpath(e)) for e in expressions]
    accel.enable()
    assert [describe(body.xpath(e)) for e in expressions] == plain


def test_later_namespace_registration_is_honoured(accelerated):
    from docx.oxml import ns
    body = Document().element.body
    with pytest.raises(etree.XPathEvalError):
        body.xpath("zz:thing")
    ns.nsmap["zz"] = "urn:example:zz"
    try:
        assert body.xpath("zz:thing") == []
    finally:
        del ns.nsmap["zz"]


def test_render_is_identical_with_accel(tmp_path, accelerated):
    doc = Document()
    doc.add_paragraph("{{ a }}")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "{{ b }}"
    path = tmp_path / "t.docx"
    doc.save(path)
    outputs = []
    for enabled in (True, False):
        (accel.enable if enabled else accel.disable)()
        tpl = docxtpl.DocxTemplate(path)
        tpl.render({"a": "x", "b": "y"})
        outputs.append(etree.tostring(tpl.docx.element, method="c14n"))
    assert outputs[0] == outputs[1]
