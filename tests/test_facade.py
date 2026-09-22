"""Behaviour of the compatibility layer that differs in *mechanism* from
docxtpl 0.20.1 (caches, native kernels, body swap) and must not differ in
observable results."""
import io
from zipfile import ZipFile

from docx import Document
import jinja2
from jinja2 import Environment
from lxml import etree
import pytest

import docxtpl
from docxtpl import DocxTemplate, _reference


@pytest.fixture(autouse=True)
def fresh_state():
    docxtpl.cache_clear()
    docxtpl.reset_stats()
    yield
    docxtpl.set_cache_budget(32 * 1024 * 1024)


def make_docx(tmp_path, *paragraphs, name="template.docx"):
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    path = tmp_path / name
    doc.save(path)
    return path


def body_text(tpl):
    stream = io.BytesIO()
    tpl.save(stream)
    with ZipFile(stream) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    return "|".join(root.xpath("//*[local-name()='t']/text()"))


def counters():
    return docxtpl.stats()["counters"]


# -- variable scan cache ------------------------------------------------------

def test_scan_cache_returns_independent_sets(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ a }} {{ b }}"))
    first = tpl.get_undeclared_template_variables()
    first.add("polluted")
    assert tpl.get_undeclared_template_variables() == {"a", "b"}
    assert tpl.get_undeclared_template_variables(context={"a": 1}) == {"b"}
    assert counters()["cache_variables_hit"] == 2


def test_scan_cache_is_keyed_by_content_not_path(tmp_path):
    path = make_docx(tmp_path, "{{ before }}")
    assert DocxTemplate(path).get_undeclared_template_variables() == {"before"}
    make_docx(tmp_path, "{{ after }}")  # same path, new content
    assert DocxTemplate(path).get_undeclared_template_variables() == {"after"}


def test_scan_reads_template_file_not_live_document(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ original }}"))
    tpl.init_docx()
    tpl.docx.add_paragraph("{{ added_later }}")
    assert tpl.get_undeclared_template_variables() == {"original"}
    tpl.render({"original": "x", "added_later": "y"})
    assert tpl.get_undeclared_template_variables() == {"original"}
    assert body_text(tpl) == "x|y"


def test_scan_depends_on_registered_filter_names(tmp_path):
    path = make_docx(tmp_path, "{{ value|custom }}")
    env = Environment()
    env.filters["custom"] = str
    assert DocxTemplate(path).get_undeclared_template_variables(jinja_env=env) == {"value"}
    with pytest.raises(jinja2.TemplateAssertionError):
        DocxTemplate(path).get_undeclared_template_variables(jinja_env=Environment())
    # a mutated environment is re-evaluated, not remembered by identity
    del env.filters["custom"]
    with pytest.raises(jinja2.TemplateAssertionError):
        DocxTemplate(path).get_undeclared_template_variables(jinja_env=env)


def test_scan_stream_position_matches_uncached_run(tmp_path):
    data = make_docx(tmp_path, "{{ a }}").read_bytes()
    docxtpl.set_cache_budget(0)
    stream = io.BytesIO(data)
    stream.seek(7)
    DocxTemplate(stream).get_undeclared_template_variables()
    expected = stream.tell()
    docxtpl.set_cache_budget(1 << 20)
    for _ in range(3):  # miss, then hits
        stream = io.BytesIO(data)
        stream.seek(7)
        assert DocxTemplate(stream).get_undeclared_template_variables() == {"a"}
        assert stream.tell() == expected
    assert counters()["cache_variables_hit"] == 2


def test_customised_patch_xml_is_never_cached(tmp_path):
    calls = []

    class Custom(DocxTemplate):
        def patch_xml(self, src_xml):
            calls.append(1)
            return super().patch_xml(src_xml).replace("{{ a }}", "{{ renamed }}")

    path = make_docx(tmp_path, "{{ a }}")
    assert Custom(path).get_undeclared_template_variables() == {"renamed"}
    assert Custom(path).get_undeclared_template_variables() == {"renamed"}
    assert len(calls) == 2
    assert DocxTemplate(path).get_undeclared_template_variables() == {"a"}


def test_extension_environment_uses_source_layer_only(tmp_path):
    path = make_docx(tmp_path, "{{ a }}")
    env = Environment(extensions=["jinja2.ext.do"])
    for _ in range(2):
        assert DocxTemplate(path).get_undeclared_template_variables(jinja_env=env) == {"a"}
    assert counters()["cache_source_hit"] == 1
    assert "cache_variables_hit" not in counters()


def test_cache_can_be_disabled(tmp_path):
    docxtpl.set_cache_budget(0)
    tpl = DocxTemplate(make_docx(tmp_path, "{{ a }}"))
    assert tpl.get_undeclared_template_variables() == {"a"}
    tpl.render({"a": "x"})
    assert not any(name.startswith("cache_") for name in counters())
    assert docxtpl.cache_info()["used_bytes"] == 0


# -- compiled template reuse --------------------------------------------------

def stamp(value):
    stamp.calls += 1
    return "%s#%d" % (value, stamp.calls)


def test_constant_folded_custom_filter_runs_on_every_render(tmp_path):
    # upstream compiles per render, so a custom filter over constants is
    # re-evaluated each time; a reused code object would freeze its result
    stamp.calls = 0
    path = make_docx(tmp_path, "{{ 'k'|stamp }}")
    results = []
    for _ in range(3):
        env = Environment()
        env.filters["stamp"] = stamp
        tpl = DocxTemplate(path)
        tpl.render({}, jinja_env=env)
        results.append(body_text(tpl))
    assert results == ["k#1", "k#2", "k#3"]


def shout(value):
    return str(value).upper()


def whisper(value):
    return str(value).lower()


def test_compiled_code_is_shared_between_equivalent_environments(tmp_path):
    path = make_docx(tmp_path, "{{ name|voice }}")
    for expected in ("ADA", "ADA"):
        env = Environment()
        env.filters["voice"] = shout
        tpl = DocxTemplate(path)
        tpl.render({"name": "Ada"}, jinja_env=env)
        assert body_text(tpl) == expected
    assert counters()["cache_code_hit"] >= 1
    # same names, different function: not equivalent
    env = Environment()
    env.filters["voice"] = whisper
    tpl = DocxTemplate(path)
    tpl.render({"name": "Ada"}, jinja_env=env)
    assert body_text(tpl) == "ada"


def test_closure_filters_are_never_captured_by_the_cache(tmp_path):
    path = make_docx(tmp_path, "{{ name|tag }}")
    for suffix in ("-1", "-2"):
        env = Environment()
        env.filters["tag"] = lambda value, suffix=suffix: value + suffix  # default arg: no closure
        env.filters["wrap"] = (lambda s: (lambda value: value + s))(suffix)  # closure
        tpl = DocxTemplate(path)
        tpl.render({"name": "n"}, jinja_env=env)
        assert body_text(tpl) == "n" + suffix
    assert "cache_code_hit" not in counters()
    assert docxtpl.cache_info()["layers"].get("code") is None


def test_autoescape_and_undefined_are_part_of_the_fingerprint(tmp_path):
    path = make_docx(tmp_path, "{{ v }}{{ missing }}")
    seen = []
    for autoescape in (False, True, False, True):
        tpl = DocxTemplate(path)
        tpl.render({"v": "a<b"}, jinja_env=Environment(autoescape=autoescape))
        seen.append(body_text(tpl))
    # unescaped "<b" is malformed XML that the recovering parser drops (as upstream)
    assert seen == ["a", "a<b", "a", "a<b"]
    with pytest.raises(jinja2.UndefinedError):
        DocxTemplate(path).render({"v": "x"}, jinja_env=Environment(undefined=jinja2.StrictUndefined))


def test_template_errors_are_canonical(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{% if %}"))
    with pytest.raises(jinja2.TemplateSyntaxError) as info:
        tpl.render({}, jinja_env=Environment())
    assert info.value.source is not None
    assert hasattr(info.value, "docx_context")


# -- render lifecycle ---------------------------------------------------------

def test_render_keeps_document_part_and_stale_views_like_upstream(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ a }}", "static"))
    document = tpl.get_docx()
    part = document.part
    held = document.paragraphs[1]  # also caches document._body, as the application does
    tpl.render({"a": "x"})
    assert counters()["map_tree_swap"] == 1
    assert tpl.docx is document and document.part is part
    assert part._element is document._element
    # upstream quirk kept on purpose: the cached body view is stale but complete
    assert [p.text for p in document.paragraphs] == ["{{ a }}", "static"]
    assert held.text == "static"
    assert [t.text for t in document._element.body.iter("{*}t")] == ["x", "static"]
    assert body_text(tpl) == "x|static"


def test_shared_root_element_takes_the_upstream_path(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ a }}"))
    root = tpl.get_docx()._element  # a third reference: identity must survive
    tpl.render({"a": "x"})
    assert counters()["map_tree_replace"] == 1
    assert tpl.docx._element is root
    assert [t.text for t in root.iter("{*}t")] == ["x"]


def test_root_level_siblings_and_attributes_survive(tmp_path):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    path = make_docx(tmp_path, "{{ a }}")
    tpl = DocxTemplate(path)
    root = tpl.get_docx()._element
    root.insert(0, parse_xml('<w:background %s w:color="FFEEDD"/>' % nsdecls("w")))
    root.set("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable", "w14")
    del root
    tpl.render({"a": "x"})
    assert counters()["map_tree_swap"] == 1
    with ZipFile(io.BytesIO(save_bytes(tpl))) as archive:
        saved = etree.fromstring(archive.read("word/document.xml"))
    assert [etree.QName(child).localname for child in saved] == ["background", "body"]
    assert saved[0].get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}color") == "FFEEDD"
    assert saved.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable") == "w14"


def test_render_is_not_reentrant_on_one_instance(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ 1|again }}"))
    env = Environment()
    errors = []

    def again(value):
        try:
            tpl.render({}, jinja_env=env)
        except RuntimeError as exc:
            errors.append(str(exc))
        return value

    env.filters["again"] = lambda value: again(value)
    tpl.render({}, jinja_env=env)
    assert errors and "already running" in errors[0]
    tpl.render({}, jinja_env=env)  # the lock is released afterwards


def test_failed_render_releases_the_instance(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ a.b }}"))
    with pytest.raises(jinja2.UndefinedError):
        tpl.render({}, jinja_env=Environment(undefined=jinja2.StrictUndefined))
    tpl.render({"a": {"b": "ok"}})
    assert body_text(tpl) == "ok"


# -- native / reference routing ----------------------------------------------

def test_lone_surrogates_take_the_reference_path(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "x"))
    text = "<w:t>{{ a\ud800 }}</w:t>"
    assert tpl.patch_xml(text) == _reference.patch_xml(text)
    assert counters()["patch_xml_reference"] == 1


def test_customised_resolve_listing_is_honoured(tmp_path):
    class Custom(DocxTemplate):
        def resolve_listing(self, xml):
            return super().resolve_listing(xml).replace("plain", "custom")

    tpl = Custom(make_docx(tmp_path, "{{ v }}"))
    tpl.render({"v": "plain\ttext"})
    assert body_text(tpl) == "custom|text"
    assert counters()["post_render_reference"] >= 1


def test_listing_characters(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ v }}"))
    tpl.render({"v": docxtpl.Listing("a\nb\tc\x07d\x0ce")})
    document = Document(io.BytesIO(save_bytes(tpl)))
    assert [p.text for p in document.paragraphs][:2] == ["a\nb\tc", "d"]


def save_bytes(tpl):
    stream = io.BytesIO()
    tpl.save(stream)
    return stream.getvalue()


def test_stage_timings_are_opt_in(tmp_path):
    tpl = DocxTemplate(make_docx(tmp_path, "{{ v }}"))
    tpl.render({"v": 1})
    assert docxtpl.stats()["timings"] == {}
    docxtpl.enable_timings()
    try:
        tpl.render({"v": 2})
        assert {"render", "patch_xml", "map_tree"} <= set(docxtpl.stats()["timings"])
    finally:
        docxtpl.enable_timings(False)


def test_subclass_without_super_init(tmp_path):
    class Minimal(DocxTemplate):
        def __init__(self, template_file):  # mirrors upstream attributes by hand
            self.template_file = template_file
            self.reset_replacements()
            self.docx = None
            self.is_rendered = self.is_saved = self.allow_missing_pics = False

    tpl = Minimal(make_docx(tmp_path, "{{ a }}"))
    tpl.render({"a": "x"})
    assert body_text(tpl) == "x"


def test_pure_python_control_group_is_equivalent(tmp_path):
    import os
    import subprocess
    import sys
    path = make_docx(tmp_path, "{{ a }}", "{%p if b %}", "shown\t{{ b }}", "{%p endif %}")
    script = (
        "import io, sys, docxtpl\n"
        "tpl = docxtpl.DocxTemplate(sys.argv[1]); tpl.render({'a': 'x', 'b': 'y'})\n"
        "tpl.save(sys.argv[2]); print(sorted(docxtpl.stats()['counters']))\n"
    )
    outputs = []
    for native in ("1", "0"):
        target = tmp_path / ("native%s.docx" % native)
        result = subprocess.run([sys.executable, "-W", "ignore", "-c", script, str(path), str(target)],
                                env=dict(os.environ, DOCXTPL_NATIVE=native), capture_output=True, text=True, check=True)
        assert ("patch_xml_native" in result.stdout) == (native == "1")
        with ZipFile(target) as archive:
            outputs.append(archive.read("word/document.xml"))
    assert outputs[0] == outputs[1]


def test_public_names_are_declared_for_type_checkers():
    # py.typed + missing __all__ makes pyright reject `from docxtpl import DocxTemplate`
    import docxtpl

    for name in ("DocxTemplate", "InlineImage", "RichText", "R", "RichTextParagraph",
                 "RP", "Listing", "Subdoc"):
        assert name in docxtpl.__all__ and hasattr(docxtpl, name)
    assert all(hasattr(docxtpl, name) for name in docxtpl.__all__)


# -- on-disk compiled-code cache ----------------------------------------------


@pytest.fixture()
def code_cache_dir(tmp_path):
    directory = tmp_path / "code-cache"
    docxtpl.configure_code_cache(directory)
    yield directory
    docxtpl.configure_code_cache(None)


def render_with_filter(path, function):
    env = Environment()
    env.filters["voice"] = function
    tpl = DocxTemplate(path)
    tpl.render({"name": "Ada"}, jinja_env=env)
    return body_text(tpl)


def test_code_cache_is_off_unless_configured(tmp_path):
    path = make_docx(tmp_path, "{{ name }}")
    DocxTemplate(path).render({"name": "x"})
    assert docxtpl.code_cache_info()["directory"] is None
    assert not any(k.startswith("code_cache") for k in counters())


def test_code_cache_survives_a_cleared_process_cache(tmp_path, code_cache_dir):
    path = make_docx(tmp_path, "{{ name|voice }}")
    assert render_with_filter(path, shout) == "ADA"
    assert counters()["code_cache_store"] == 1
    assert docxtpl.code_cache_info()["entries"] == 1
    docxtpl.cache_clear()
    docxtpl.reset_stats()
    assert render_with_filter(path, shout) == "ADA"
    assert counters()["code_cache_hit"] == 1
    assert counters()["jinja_compile_reused"] == 1
    assert "code_cache_store" not in counters()


def test_code_cache_key_covers_filter_code_not_identity(tmp_path, code_cache_dir):
    path = make_docx(tmp_path, "{{ name|voice }}")
    assert render_with_filter(path, shout) == "ADA"
    docxtpl.cache_clear()
    docxtpl.reset_stats()
    assert render_with_filter(path, whisper) == "ada"
    assert counters()["code_cache_miss"] == 1
    assert docxtpl.code_cache_info()["entries"] == 2


def test_code_cache_is_shared_with_a_fresh_interpreter(tmp_path, code_cache_dir):
    path = make_docx(tmp_path, "{{ name|upper }} {{ 2 + 2 }}")
    DocxTemplate(path).render({"name": "Ada"})
    assert counters()["code_cache_store"] == 1
    import os
    import subprocess
    import sys

    script = (
        "import docxtpl, sys; from docxtpl import DocxTemplate;"
        "tpl = DocxTemplate(sys.argv[1]); tpl.render({'name': 'Ada'});"
        "import io, zipfile; from lxml import etree; s = io.BytesIO(); tpl.save(s);"
        "print(''.join(etree.fromstring(zipfile.ZipFile(s).read('word/document.xml'))"
        ".xpath('//*[local-name()=\"t\"]/text()')));"
        "print(docxtpl.stats()['counters'])"
    )
    env = dict(os.environ, DOCXTPL_CODE_CACHE_DIR=str(code_cache_dir))
    out = subprocess.run(
        [sys.executable, "-c", script, str(path)], env=env, check=True, capture_output=True, text=True
    ).stdout
    assert out.splitlines()[0] == "ADA 4"
    assert "'code_cache_hit': 1" in out
    assert "code_cache_store" not in out


def test_corrupt_code_cache_entry_is_replaced(tmp_path, code_cache_dir):
    path = make_docx(tmp_path, "{{ name }}")
    DocxTemplate(path).render({"name": "x"})
    (entry,) = list(code_cache_dir.iterdir())
    entry.write_bytes(b"garbage")
    docxtpl.cache_clear()
    docxtpl.reset_stats()
    tpl = DocxTemplate(path)
    tpl.render({"name": "y"})
    assert body_text(tpl) == "y"
    assert counters()["code_cache_error"] == 1
    assert counters()["code_cache_store"] == 1
    assert entry.read_bytes() != b"garbage"


def test_uncacheable_constant_folding_never_reaches_disk(tmp_path, code_cache_dir):
    path = make_docx(tmp_path, "{{ 'ada'|voice }}")  # constant input: folded at compile time
    assert render_with_filter(path, shout) == "ADA"
    assert render_with_filter(path, whisper) == "ada"
    assert docxtpl.code_cache_info()["entries"] == 0


def test_closure_filters_never_reach_disk(tmp_path, code_cache_dir):
    path = make_docx(tmp_path, "{{ name|tag }}")
    env = Environment()
    env.filters["tag"] = (lambda s: (lambda value: value + s))("-1")
    tpl = DocxTemplate(path)
    tpl.render({"name": "n"}, jinja_env=env)
    assert body_text(tpl) == "n-1"
    assert docxtpl.code_cache_info()["entries"] == 0


def test_code_cache_prunes_oldest_entries(tmp_path):
    from docxtpl import _codecache

    directory = tmp_path / "bounded"
    docxtpl.configure_code_cache(directory, max_entries=3)
    _codecache._puts = 0  # pruning runs every _PRUNE_EVERY stores, process-wide
    try:
        for index in range(_codecache._PRUNE_EVERY):
            DocxTemplate(make_docx(tmp_path, "{{ v }} %d" % index, name="t%d.docx" % index)).render({"v": 1})
        assert docxtpl.code_cache_info()["entries"] == 3
    finally:
        docxtpl.configure_code_cache(None)
