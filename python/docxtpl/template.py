# -*- coding: utf-8 -*-
"""
Created : 2015-03-12

@author: Eric Lapouyade
"""

from os import PathLike
from typing import Any, Optional, IO, Union, Dict, Set
from .subdoc import Subdoc
import functools
import io
from lxml import etree
from docx import Document
from docx.opc.oxml import parse_xml
from docx.opc.part import XmlPart
import docx.oxml.ns
from docx.opc.constants import RELATIONSHIP_TYPE as REL_TYPE
from jinja2 import Environment, Template, meta
from jinja2.exceptions import TemplateError

try:
    from html import escape  # noqa: F401
except ImportError:
    # cgi.escape is deprecated in python 3.7
    from cgi import escape  # noqa: F401
import re
import sys
import binascii
import copy
import hashlib
import os
import zipfile

from . import _reference, _stats
from . import _jinja
from ._cache import cache as _cache, str_nbytes


if os.environ.get("DOCXTPL_NATIVE", "1") == "0":
    # Control group for benchmarks: same algorithms, reference text kernels.
    class _native(object):
        @staticmethod
        def _declined(text):
            raise TypeError("native kernels disabled by DOCXTPL_NATIVE=0")

        patch_xml = pre_render = post_render = _declined
else:
    from . import _native


def _element_refs(holder):
    return sys.getrefcount(holder._element)


def _calibrate_unshared_refs():
    """``_element_refs()`` of an element referenced by exactly two holders
    (as the root is by its ``Document`` and ``DocumentPart``); measured, not
    assumed, because the bookkeeping differs between CPython versions."""
    if not hasattr(sys, "getrefcount"):
        return None

    class Holder(object):
        pass

    first, second = Holder(), Holder()
    first._element = second._element = etree.Element("probe")
    return _element_refs(first)


_UNSHARED_ROOT_REFS = _calibrate_unshared_refs()


def _default_environment():
    """The spontaneous environment ``jinja2.Template(source)`` compiles in."""
    return Template("").environment


class DocxTemplate(object):
    """Class for managing docx files as they were jinja2 templates"""

    HEADER_URI = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header"
    )
    FOOTER_URI = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footer"
    )

    def __init__(self, template_file: Union[IO[bytes], str, PathLike]) -> None:
        self.template_file = template_file
        self.reset_replacements()
        self.docx = None
        self.is_rendered = False
        self.is_saved = False
        self.allow_missing_pics = False

    def init_docx(self, reload: bool = True):
        if not self.docx or (self.is_rendered and reload):
            self.docx = Document(self.template_file)
            self.is_rendered = False

    def render_init(self):
        self.init_docx()
        self.pic_map = {}
        self.current_rendering_part = None
        self.docx_ids_index = 1000
        self.is_saved = False

    def __getattr__(self, name):
        return getattr(self.docx, name)

    def xml_to_string(self, xml, encoding="unicode"):
        # Be careful : pretty_print MUST be set to False, otherwise patch_xml()
        # won't work properly
        return etree.tostring(xml, encoding="unicode", pretty_print=False)

    def get_docx(self):
        self.init_docx()
        return self.docx

    def get_xml(self):
        return self.xml_to_string(self.docx._element.body)

    def write_xml(self, filename):
        with open(filename, "w") as fh:
            fh.write(self.get_xml())

    def patch_xml(self, src_xml):
        """Make a lots of cleaning to have a raw xml understandable by jinja2 :
        strip all unnecessary xml tags, manage table cell background color and colspan,
        unescape html entities, etc..."""
        with _stats.stage("patch_xml"):
            try:
                patched = _native.patch_xml(src_xml)
            except (TypeError, UnicodeEncodeError):
                # not a str, or lone surrogates that cannot cross as UTF-8
                patched = None
            if patched is None:
                _stats.count("patch_xml_reference")
                return _reference.patch_xml(src_xml)
            _stats.count("patch_xml_native")
            return patched

    def render_xml_part(self, src_xml, part, context, jinja_env=None):
        try:
            src_xml = _native.pre_render(src_xml) or src_xml
        except (TypeError, UnicodeEncodeError):
            src_xml = re.sub(r"<w:p([ >])", r"\n<w:p\1", src_xml)
        try:
            self.current_rendering_part = part
            if jinja_env:
                template = _jinja.template_from_string(jinja_env, src_xml)
            else:
                # Template(src) compiles in Jinja2's shared default environment
                template = _jinja.template_from_string(_default_environment(), src_xml)
            dst_xml = template.render(context)
        except TemplateError as exc:
            if hasattr(exc, "lineno") and exc.lineno is not None:
                line_number = max(exc.lineno - 4, 0)
                exc.docx_context = map(
                    lambda x: re.sub(r"<[^>]+>", "", x),
                    src_xml.splitlines()[line_number: (line_number + 7)],  # fmt: skip
                )

            raise exc
        with _stats.stage("post_render"):
            return self._post_render(dst_xml)

    def _post_render(self, dst_xml):
        # One native call unless resolve_listing() is customised or the text
        # cannot cross the FFI boundary (lone surrogates from the context).
        if type(self).resolve_listing is DocxTemplate.resolve_listing:
            try:
                rendered = _native.post_render(dst_xml)
                _stats.count("post_render_native")
                return dst_xml if rendered is None else rendered
            except (TypeError, UnicodeEncodeError):
                pass
        _stats.count("post_render_reference")
        dst_xml = re.sub(r"\n<w:p([ >])", r"<w:p\1", dst_xml)
        dst_xml = (
            dst_xml.replace("{_{", "{{")
            .replace("}_}", "}}")
            .replace("{_%", "{%")
            .replace("%_}", "%}")
        )
        dst_xml = self.resolve_listing(dst_xml)
        return dst_xml

    def render_properties(
        self, context: Dict[str, Any], jinja_env: Optional[Environment] = None
    ) -> None:
        # List of string attributes of docx.opc.coreprops.CoreProperties which are strings.
        # It seems that some attributes cannot be written as strings. Those are commented out.
        properties = [
            "author",
            # 'category',
            "comments",
            # 'content_status',
            "identifier",
            # 'keywords',
            "language",
            # 'last_modified_by',
            "subject",
            "title",
            # 'version',
        ]
        if jinja_env is None:
            jinja_env = Environment()

        for prop in properties:
            initial = getattr(self.docx.core_properties, prop)
            template = jinja_env.from_string(initial)
            rendered = template.render(context)
            setattr(self.docx.core_properties, prop, rendered)

    def render_footnotes(
        self, context: Dict[str, Any], jinja_env: Optional[Environment] = None
    ) -> None:
        if jinja_env is None:
            jinja_env = Environment()

        for section in self.docx.sections:
            for part in section.part.package.parts:
                if part.content_type == (
                    "application/vnd.openxmlformats-officedocument"
                    ".wordprocessingml.footnotes+xml"
                ):
                    xml = self.patch_xml(
                        part.blob.decode("utf-8")
                        if isinstance(part.blob, bytes)
                        else part.blob
                    )
                    xml = self.render_xml_part(xml, part, context, jinja_env)
                    part._blob = xml.encode("utf-8")

    def resolve_listing(self, xml):
        return _reference.resolve_listing(xml)

    def build_xml(self, context, jinja_env=None):
        xml = self.get_xml()
        xml = self.patch_xml(xml)
        xml = self.render_xml_part(xml, self.docx._part, context, jinja_env)
        return xml

    def map_tree(self, tree):
        with _stats.stage("map_tree"):
            if not self._swap_root(tree):
                _stats.count("map_tree_replace")
                root = self.docx._element
                body = root.body
                root.replace(body, tree)

    def _swap_root(self, tree):
        """Install the rendered body without unlinking the old one.

        ``root.replace(body, tree)`` makes lxml re-home every node of the
        outgoing body; lxml's namespace cache never hits for namespaces
        declared on the root, so this is quadratic in the body size and
        dominates upstream ``render()`` on large templates. Instead, a twin
        of the root element is created, the new body is attached to it (the
        cheap direction) and the python-docx ``Document`` and ``DocumentPart``
        are pointed at the twin. The previous tree is left untouched, so
        anything obtained before ``render()`` (a cached ``_Body``, tables,
        paragraphs) stays complete and stale, exactly as upstream.

        The one observable difference would be the identity of the root
        element, so the swap only happens when nothing but the document, its
        part and this frame reference it. Returns False to request the
        upstream path."""
        document = self.docx
        if getattr(document._part, "_element", None) is not document._element:
            return False
        if _element_refs(document) != _UNSHARED_ROOT_REFS:
            return False
        root = document._element
        if root.getprevious() is not None or root.getnext() is not None:
            return False
        body = root.body
        if body is None or body.getparent() is not root:
            return False

        twin = root.makeelement(root.tag, attrib=dict(root.attrib), nsmap=root.nsmap)
        twin.text = root.text
        for child in root:
            if child is body:
                node = tree
            else:
                node = copy.deepcopy(child)
            node.tail = child.tail
            twin.append(node)

        document._element = twin
        document._part._element = twin
        _stats.count("map_tree_swap")
        return True

    def get_headers_footers(self, uri):
        for relKey, val in self.docx._part.rels.items():
            if (val.reltype == uri) and (val.target_part.blob):
                yield relKey, val.target_part

    def get_part_xml(self, part):
        return self.xml_to_string(parse_xml(part.blob))

    def get_headers_footers_encoding(self, xml):
        m = re.match(r'<\?xml[^\?]+\bencoding="([^"]+)"', xml, re.I)
        if m:
            return m.group(1)
        return "utf-8"

    def build_headers_footers_xml(self, context, uri, jinja_env=None):
        for relKey, part in self.get_headers_footers(uri):
            xml = self.get_part_xml(part)
            encoding = self.get_headers_footers_encoding(xml)
            xml = self.patch_xml(xml)
            xml = self.render_xml_part(xml, part, context, jinja_env)
            yield relKey, xml.encode(encoding)

    def map_headers_footers_xml(self, relKey, xml):
        part = self.docx._part.rels[relKey].target_part
        new_part = XmlPart.load(part.partname, part.content_type, xml, part.package)
        for rId, rel in part.rels.items():
            new_part.load_rel(rel.reltype, rel._target, rel.rId, rel.is_external)
        self.docx._part.rels[relKey]._target = new_part

    def render(
        self,
        context: Dict[str, Any],
        jinja_env: Optional[Environment] = None,
        autoescape: bool = False,
    ) -> None:
        # A plain flag read from __dict__ also covers subclasses whose
        # __init__ does not call ours (__getattr__ proxies to the document).
        if self.__dict__.get("_rendering"):
            raise RuntimeError(
                "DocxTemplate.render() is already running on this instance; "
                "use one DocxTemplate per thread"
            )
        self._rendering = True
        try:
            with _stats.stage("render"):
                self._render(context, jinja_env, autoescape)
        finally:
            self._rendering = False

    def _render(self, context, jinja_env, autoescape):
        # init template working attributes
        self.render_init()

        if autoescape:
            if not jinja_env:
                jinja_env = Environment(autoescape=autoescape)
            else:
                jinja_env.autoescape = autoescape

        # Body
        xml_src = self.build_xml(context, jinja_env)

        # fix tables if needed
        tree = self.fix_tables(xml_src)

        # fix docPr ID's
        self.fix_docpr_ids(tree)

        # Replace body xml tree
        self.map_tree(tree)

        # Headers
        headers = self.build_headers_footers_xml(context, self.HEADER_URI, jinja_env)
        for relKey, xml in headers:
            self.map_headers_footers_xml(relKey, xml)

        # Footers
        footers = self.build_headers_footers_xml(context, self.FOOTER_URI, jinja_env)
        for relKey, xml in footers:
            self.map_headers_footers_xml(relKey, xml)

        self.render_properties(context, jinja_env)

        self.render_footnotes(context, jinja_env)

        # set rendered flag
        self.is_rendered = True

    # using of TC tag in for cycle can cause that count of columns does not
    # correspond to real count of columns in row. This function is able to fix it.
    def fix_tables(self, xml):
        parser = etree.XMLParser(recover=True)
        tree = etree.fromstring(xml, parser=parser)
        # get namespace
        ns = "{" + tree.nsmap["w"] + "}"
        # walk trough xml and find table
        for t in tree.iter(ns + "tbl"):
            tblGrid = t.find(ns + "tblGrid")
            columns = tblGrid.findall(ns + "gridCol")
            to_add = 0
            # walk trough all rows and try to find if there is higher cell count
            for r in t.iter(ns + "tr"):
                cells = r.findall(ns + "tc")
                if (len(columns) + to_add) < len(cells):
                    to_add = len(cells) - len(columns)
            # is necessary to add columns?
            if to_add > 0:
                # at first, calculate width of table according to columns
                # (we want to preserve it)
                width = 0.0
                new_average = None
                for c in columns:
                    if not c.get(ns + "w") is None:
                        width += float(c.get(ns + "w"))
                # try to keep proportion of table
                if width > 0:
                    old_average = width / len(columns)
                    new_average = width / (len(columns) + to_add)
                    # scale the old columns
                    for c in columns:
                        c.set(
                            ns + "w",
                            str(
                                int(float(c.get(ns + "w")) * new_average / old_average)
                            ),
                        )
                    # add new columns
                    for i in range(to_add):
                        etree.SubElement(
                            tblGrid, ns + "gridCol", {ns + "w": str(int(new_average))}
                        )

            # Refetch columns after columns addition.
            columns = tblGrid.findall(ns + "gridCol")
            columns_len = len(columns)

            cells_len_max = 0

            def get_cell_len(total, cell):
                tc_pr = cell.find(ns + "tcPr")
                grid_span = None if tc_pr is None else tc_pr.find(ns + "gridSpan")

                if grid_span is not None:
                    return total + int(grid_span.get(ns + "val"))

                return total + 1

            # Calculate max of table cells to compare with `gridCol`.
            for r in t.iter(ns + "tr"):
                cells = r.findall(ns + "tc")
                cells_len = functools.reduce(get_cell_len, cells, 0)
                cells_len_max = max(cells_len_max, cells_len)

            to_remove = columns_len - cells_len_max

            # If after the loop, there're less columns, than
            # originally was, remove extra `gridCol` declarations.
            if to_remove > 0:
                # Have to keep track of the removed width to scale the
                # table back to its original width.
                removed_width = 0.0

                for c in columns[-to_remove:]:
                    removed_width += float(c.get(ns + "w"))

                    tblGrid.remove(c)

                columns_left = tblGrid.findall(ns + "gridCol")

                # Distribute `removed_width` across all columns that has
                # left after extras removal.
                extra_space = 0
                if len(columns_left) > 0:
                    extra_space = removed_width / len(columns_left)
                    extra_space = int(extra_space)

                for c in columns_left:
                    c.set(ns + "w", str(int(float(c.get(ns + "w")) + extra_space)))

        return tree

    def fix_docpr_ids(self, tree):
        # some Ids may have some collisions : so renumbering all of them :
        for elt in tree.xpath("//wp:docPr", namespaces=docx.oxml.ns.nsmap):
            self.docx_ids_index += 1
            elt.attrib["id"] = str(self.docx_ids_index)

    def new_subdoc(self, docpath=None):
        self.init_docx()
        return Subdoc(self, docpath)

    @staticmethod
    def get_file_crc(file_obj):
        if hasattr(file_obj, "read"):
            buf = file_obj.read()
        else:
            with open(file_obj, "rb") as fh:
                buf = fh.read()

        crc = binascii.crc32(buf) & 0xFFFFFFFF
        return crc

    def replace_media(self, src_file, dst_file):
        """Replace one media by another one into a docx

        This has been done mainly because it is not possible to add images in
        docx header/footer.
        With this function, put a dummy picture in your header/footer,
        then specify it with its replacement in this function using the file path
        or file-like objects.

        Syntax: tpl.replace_media('dummy_media_to_replace.png','media_to_paste.jpg')
            -- or --
                tpl.replace_media(io.BytesIO(image_stream), io.BytesIO(new_image_stream))

        Note: for images, the aspect ratio will be the same as the replaced image

        Note2: it is important to have the source media file as it is required
                to calculate its CRC to find them in the docx
        """

        crc = self.get_file_crc(src_file)
        if hasattr(dst_file, "read"):
            self.crc_to_new_media[crc] = dst_file.read()
        else:
            with open(dst_file, "rb") as fh:
                self.crc_to_new_media[crc] = fh.read()

    def replace_pic(self, embedded_file, dst_file):
        """Replace embedded picture with original-name given by embedded_file.
           (give only the file basename, not the full path)
           The new picture is given by dst_file (either a filename or a file-like
           object)

        Notes:
            1) embedded_file and dst_file must have the same extension/format
               in case dst_file is a file-like object, no check is done on
               format compatibility
            2) the aspect ratio will be the same as the replaced image
            3) There is no need to keep the original file (this is not the case
               for replace_embedded and replace_media)
        """

        if hasattr(dst_file, "read"):
            # NOTE: file extension not checked
            self.pics_to_replace[embedded_file] = dst_file.read()
        else:
            with open(dst_file, "rb") as fh:
                self.pics_to_replace[embedded_file] = fh.read()

    def replace_embedded(self, src_file, dst_file):
        """Replace one embedded object by another one into a docx

        This has been done mainly because it is not possible to add images
        in docx header/footer.
        With this function, put a dummy picture in your header/footer,
        then specify it with its replacement in this function

        Syntax: tpl.replace_embedded('dummy_doc.docx','doc_to_paste.docx')

        Note2 : it is important to have the source file as it is required to
                calculate its CRC to find them in the docx
        """
        with open(dst_file, "rb") as fh:
            crc = self.get_file_crc(src_file)
            self.crc_to_new_embedded[crc] = fh.read()

    def replace_zipname(self, zipname, dst_file):
        """Replace one file in the docx file

        First note that a MSWord .docx file is in fact a zip file.

        This method can be used to replace document embedded in the docx template.

        Some embedded document may have been modified by MSWord while saving
        the template : thus replace_embedded() cannot be used as CRC is not the
        same as the original file.

        This method works for embedded MSWord file like Excel or PowerPoint file,
        but won't work for others like PDF, Python or even Text files :
        For these ones, MSWord generate an oleObjectNNN.bin file which is no
        use to be replaced as it is encoded.

        Syntax:

        tpl.replace_zipname(
            'word/embeddings/Feuille_Microsoft_Office_Excel1.xlsx',
            'my_excel_file.xlsx')

        The zipname is the one you can find when you open docx with WinZip,
        7zip (Windows) or unzip -l (Linux). The zipname starts with
        "word/embeddings/". Note that the file is renamed by MSWord,
        so you have to guess a little bit...
        """
        with open(dst_file, "rb") as fh:
            self.zipname_to_replace[zipname] = fh.read()

    def reset_replacements(self):
        """Reset replacement dictionaries

        This will reset data for image/embedded/zipname replacement

        This is useful when calling several times render() with different
        image/embedded/zipname replacements without re-instantiating
        DocxTemplate object.
        In this case, the right sequence for each rendering will be :
            - reset_replacements(...)
            - replace_zipname(...), replace_media(...) and/or replace_embedded(...),
            - render(...)

        If you instantiate DocxTemplate object before each render(),
        this method is useless.
        """
        self.crc_to_new_media = {}
        self.crc_to_new_embedded = {}
        self.zipname_to_replace = {}
        self.pics_to_replace = {}

    def post_processing(self, docx_file):
        if self.crc_to_new_media or self.crc_to_new_embedded or self.zipname_to_replace:

            if hasattr(docx_file, "read"):
                tmp_file = io.BytesIO()
                DocxTemplate(docx_file).save(tmp_file)
                tmp_file.seek(0)
                docx_file.seek(0)
                docx_file.truncate()
                docx_file.seek(0)

            else:
                tmp_file = "%s_docxtpl_before_replace_medias" % docx_file
                os.rename(docx_file, tmp_file)

            with zipfile.ZipFile(tmp_file) as zin:
                with zipfile.ZipFile(docx_file, "w") as zout:
                    for item in zin.infolist():
                        buf = zin.read(item.filename)
                        if item.filename in self.zipname_to_replace:
                            zout.writestr(item, self.zipname_to_replace[item.filename])
                        elif (
                            item.filename.startswith("word/media/")
                            and item.CRC in self.crc_to_new_media
                        ):
                            zout.writestr(item, self.crc_to_new_media[item.CRC])
                        elif (
                            item.filename.startswith("word/embeddings/")
                            and item.CRC in self.crc_to_new_embedded
                        ):
                            zout.writestr(item, self.crc_to_new_embedded[item.CRC])
                        else:
                            zout.writestr(item, buf)

            if not hasattr(tmp_file, "read"):
                os.remove(tmp_file)
            if hasattr(docx_file, "read"):
                docx_file.seek(0)

    def pre_processing(self):

        if self.pics_to_replace:
            self._replace_pics()

    def _replace_pics(self):
        """Replaces pictures xml tags in the docx template with pictures provided by the user"""

        replaced_pics = {key: False for key in self.pics_to_replace}

        # Main document
        part = self.docx.part
        self._replace_docx_part_pics(part, replaced_pics)

        # Header/Footer
        for relid, rel in part.rels.items():
            if rel.reltype in (REL_TYPE.HEADER, REL_TYPE.FOOTER):
                self._replace_docx_part_pics(rel.target_part, replaced_pics)

        if not self.allow_missing_pics:
            # make sure all template images defined by user were replaced
            for img_id, replaced in replaced_pics.items():
                if not replaced:
                    raise ValueError(
                        "Picture %s not found in the docx template" % img_id
                    )

    def get_pic_map(self):
        return self.pic_map

    def _replace_docx_part_pics(self, doc_part, replaced_pics):

        et = etree.fromstring(doc_part.blob)

        part_map = {}

        gds = et.xpath("//a:graphic/a:graphicData", namespaces=docx.oxml.ns.nsmap)
        for gd in gds:
            rel = None
            # Either IMAGE, CHART, SMART_ART, ...
            try:
                if gd.attrib["uri"] == docx.oxml.ns.nsmap["pic"]:
                    # Either PICTURE or LINKED_PICTURE image
                    blip = gd.xpath(
                        "pic:pic/pic:blipFill/a:blip", namespaces=docx.oxml.ns.nsmap
                    )[0]
                    dest = blip.xpath("@r:embed", namespaces=docx.oxml.ns.nsmap)
                    if len(dest) > 0:
                        rel = dest[0]
                    else:
                        continue
                else:
                    continue

                non_visual_properties = "pic:pic/pic:nvPicPr/pic:cNvPr/"
                filename = gd.xpath(
                    "%s@name" % non_visual_properties, namespaces=docx.oxml.ns.nsmap
                )[0]
                titles = gd.xpath(
                    "%s@title" % non_visual_properties, namespaces=docx.oxml.ns.nsmap
                )
                if titles:
                    title = titles[0]
                else:
                    title = ""
                descriptions = gd.xpath(
                    "%s@descr" % non_visual_properties, namespaces=docx.oxml.ns.nsmap
                )
                if descriptions:
                    description = descriptions[0]
                else:
                    description = ""

                part_map[filename] = (
                    doc_part.rels[rel].target_ref,
                    doc_part.rels[rel].target_part,
                )

                # replace data
                for img_id, img_data in self.pics_to_replace.items():
                    if img_id == filename or img_id == title or img_id == description:
                        part_map[filename][1]._blob = img_data
                        replaced_pics[img_id] = True
                        break

            # FIXME: figure out what exceptions are thrown here
            # and catch more specific exceptions
            except Exception:
                continue

        self.pic_map.update(part_map)

    def build_url_id(self, url):
        self.init_docx()
        return self.docx._part.relate_to(url, REL_TYPE.HYPERLINK, is_external=True)

    def save(self, filename: Union[IO[bytes], str, PathLike], *args, **kwargs) -> None:
        # case where save() is called without doing rendering
        # ( user wants only to replace image/embedded/zipname )
        if not self.is_saved and not self.is_rendered:
            self.docx = Document(self.template_file)
        self.pre_processing()
        self.docx.save(filename, *args, **kwargs)
        self.post_processing(filename)
        self.is_saved = True

    def get_undeclared_template_variables(
        self,
        jinja_env: Optional[Environment] = None,
        context: Optional[Dict[str, Any]] = None,
    ) -> Set[str]:
        with _stats.stage("get_undeclared_template_variables"):
            all_variables = self._undeclared_template_variables(jinja_env)

        # If context is provided, return only variables that are not in the context
        if context is not None:
            provided_variables = set(context.keys())
            return all_variables - provided_variables

        # If no context provided, return all variables (original behavior)
        return all_variables

    def _undeclared_template_variables(self, jinja_env):
        """All undeclared variables of the *template file* (never of the live,
        possibly modified or rendered, document); always a fresh ``set``."""
        if jinja_env:
            env = jinja_env
        else:
            env = Environment()

        digest, stream_start = self._template_digest()
        if digest is None:
            xml = self._template_source_xml()
            return meta.find_undeclared_variables(env.parse(xml))

        # Layer 2: the result itself, when the environment is fully described.
        fingerprint = _jinja.parse_fingerprint(env)
        if fingerprint is not None:
            cached = _cache.get(("variables", digest, fingerprint))
            if cached is not None:
                variables, stream_position = cached
                self._restore_stream(stream_position)
                return set(variables)

        # Layer 1: the patched source, valid for any environment.
        cached = _cache.get(("source", digest)) if fingerprint is None else None
        if cached is None:
            self._restore_stream(stream_start)
            xml = self._template_source_xml()
            stream_position = self._stream_position()
            if fingerprint is None:
                _cache.put(("source", digest), (xml, stream_position), str_nbytes(xml))
        else:
            xml, stream_position = cached
            self._restore_stream(stream_position)

        all_variables = meta.find_undeclared_variables(env.parse(xml))
        if fingerprint is not None:
            _cache.put(
                ("variables", digest, fingerprint),
                (frozenset(all_variables), stream_position),
                256 + 64 * len(all_variables),
            )
        return all_variables

    def _template_source_xml(self):
        # Create a temporary document to analyze the template without affecting the current state
        temp_doc = Document(self.template_file)

        # Get XML from the temporary document
        xml = self.xml_to_string(temp_doc._element.body)
        xml = self.patch_xml(xml)

        # Add headers and footers
        for uri in [self.HEADER_URI, self.FOOTER_URI]:
            for relKey, val in temp_doc._part.rels.items():
                if (val.reltype == uri) and (val.target_part.blob):
                    _xml = self.xml_to_string(parse_xml(val.target_part.blob))
                    xml += self.patch_xml(_xml)
        return xml

    def _template_digest(self):
        """``(sha256 of the template bytes, stream offset to restore)`` or
        ``(None, None)`` when analysis results must not be cached: caching
        disabled, customised ``patch_xml``/``xml_to_string``, unreadable input.
        A path or mtime is never trusted as a content identity."""
        if not _cache.enabled:
            return None, None
        for name in ("patch_xml", "xml_to_string"):
            if getattr(getattr(self, name), "__func__", None) is not getattr(DocxTemplate, name):
                return None, None
        template_file = self.template_file
        try:
            if hasattr(template_file, "read"):
                start = template_file.tell()
                template_file.seek(0)
                data = template_file.read()
                template_file.seek(start)
                if not isinstance(data, bytes):
                    return None, None
            else:
                start = None
                with open(template_file, "rb") as fh:
                    data = fh.read()
        except Exception:
            return None, None
        return hashlib.sha256(data).digest(), start

    def _stream_position(self):
        tell = getattr(self.template_file, "tell", None)
        return tell() if tell is not None and hasattr(self.template_file, "read") else None

    def _restore_stream(self, position):
        # A cache hit leaves a stream template where upstream's re-read of the
        # package would have left it.
        if position is not None:
            self.template_file.seek(position)
